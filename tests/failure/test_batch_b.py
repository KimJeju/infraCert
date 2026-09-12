"""배치 B — 조치 분류(diff)·리포트 게이트·dry-run·연결 사전검증."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from infraguard.core.models import CheckResult, HostResult, RemoteEnvironment, ScanResult
from infraguard.core.status import Status
from infraguard.orchestrator import dryrun, preflight
from infraguard.orchestrator.host_scan import HostJob, scan_host
from infraguard.reporting import html as html_report
from infraguard.reporting import xlsx as xlsx_report
from infraguard.result import diff, gate
from infraguard.rulepack import loader
from infraguard.transport.base import CleanupReport, Connection, ExecResult

PACK = Path(__file__).resolve().parents[2] / "rulepacks" / "kisa-2026"


def _r(rid: str, st: Status, **kw) -> CheckResult:
    return CheckResult(rule_id=rid, name=rid, status=st, reason="t", **kw)


def _scan(sid: str, hosts: dict[str, list[CheckResult]]) -> ScanResult:
    return ScanResult(scan_id=sid, engine_version="0", started_at=datetime.now(),
                      hosts=[HostResult(host_id=h, hostname=h, results=rs) for h, rs in hosts.items()])


# ------------------------------------------------------------------ diff
def test_diff_classifies_fixed_regressed_still_new_and_ignores_first_seen_hosts() -> None:
    base = _scan("s1", {"a": [_r("U-01", Status.FAIL), _r("U-02", Status.FAIL), _r("U-03", Status.PASS),
                              _r("U-04", Status.UNKNOWN)]})
    cur = _scan("s2", {"a": [_r("U-01", Status.PASS), _r("U-02", Status.FAIL), _r("U-03", Status.FAIL),
                             _r("U-04", Status.FAIL), _r("U-05", Status.FAIL)],
                       "b": [_r("U-01", Status.FAIL)]})            # b 는 전회에 없던 호스트
    d = diff.diff(cur, base)
    assert d[("a", "U-01")][1] == diff.FIXED
    assert d[("a", "U-02")][1] == diff.STILL_VULN
    assert d[("a", "U-03")][1] == diff.REGRESSED
    assert d[("a", "U-04")][1] == diff.NEW_VULN and d[("a", "U-05")][1] == diff.NEW_VULN
    assert d[("b", "U-01")] == (None, None)
    s = diff.summary(d)
    assert s == {diff.FIXED: 1, diff.REGRESSED: 1, diff.STILL_VULN: 1, diff.NEW_VULN: 2}
    assert diff.fix_rate(s) == 0.5
    assert diff.fix_rate({}) is None
    assert diff.diff(cur, None) == {}


def test_reports_carry_baseline_columns(tmp_path: Path) -> None:
    base = _scan("s1", {"a": [_r("U-01", Status.FAIL, evidence="x")]})
    cur = _scan("s2", {"a": [_r("U-01", Status.PASS, evidence="y")]})
    h = html_report.build(cur, tmp_path / "r.html", baseline=base).read_text(encoding="utf-8")
    assert "조치됨" in h and "조치율 100%" in h and "<th>변화</th>" in h
    x = xlsx_report.build(cur, tmp_path / "r.xlsx", baseline=base)
    from openpyxl import load_workbook
    wb = load_workbook(x)
    row = [c.value for c in wb["결과"][2]]
    assert row[-2:] == ["취약", "조치됨"]
    cells = {r[0].value: r[1].value for r in wb["요약"].iter_rows(min_row=3, max_row=16, max_col=2)}
    assert cells.get("  조치됨") == "1" and cells.get("  조치율") == "100%"
    # baseline 없이도 깨지지 않는다
    html_report.build(cur, tmp_path / "r2.html")
    xlsx_report.build(cur, tmp_path / "r2.xlsx")


# ------------------------------------------------------------------ gate
def test_gate_lists_facts_and_does_not_pass_on_open_items() -> None:
    scan = _scan("s", {"a": [_r("U-01", Status.FAIL, evidence=""), _r("U-02", Status.ERROR),
                             _r("U-03", Status.UNKNOWN), _r("U-04", Status.UNKNOWN, verdict_source="analyst")],
                       "b": []})
    scan.hosts[1].error = "boom"
    scan.rule_pack_sha256 = "a" * 64
    items = gate.run(scan, pack_sha256="b" * 64, remediation={}, db_ok=True)
    by = {i.label: i for i in items}
    assert not by["모든 호스트 진단 완료"].ok and "b" in by["모든 호스트 진단 완료"].detail
    assert not by["실행오류(ERROR) 0"].ok
    assert not by["수동확인 처리"].ok and "미판정 1건" in by["수동확인 처리"].detail
    assert not by["취약 항목 근거 존재"].ok and "a/U-01" in by["취약 항목 근거 존재"].detail
    assert not by["취약 항목 조치방법 존재"].ok and "U-01" in by["취약 항목 조치방법 존재"].detail
    assert not by["룰팩 SHA 확인"].ok and "바뀜" in by["룰팩 SHA 확인"].detail
    assert by["결과 DB 무결성"].ok
    assert not gate.passed(items)

    good = _scan("g", {"a": [_r("U-01", Status.FAIL, evidence="ev"), _r("U-02", Status.PASS)]})
    good.rule_pack_sha256 = "a" * 64
    ok = gate.run(good, pack_sha256="a" * 64, remediation={"U-01": "고쳐라"}, db_ok=True)
    assert gate.passed(ok), [i.line() for i in ok if not i.ok]


# ------------------------------------------------------------------ dry-run
def test_dryrun_enumerates_commands_per_platform_without_touching_network() -> None:
    pack = loader.load(PACK)
    lin = dryrun.plan(pack, pack.profiles["linux-native"], "linux")
    assert lin.commands and all(sh == "sh" for _, sh, _ in lin.commands)
    assert lin.violations == [] and not lin.bundles
    assert not any(rid.startswith("W-") for rid, _, _ in lin.commands)
    win = dryrun.plan(pack, pack.profiles["windows-native"], "windows")
    assert win.commands and all(sh == "powershell" for _, sh, _ in win.commands)
    jun = dryrun.plan(pack, pack.profiles["junos-native"], "junos")
    assert all(c.startswith("show") for _, _, c in jun.commands)
    mism = dryrun.plan(pack, pack.profiles["linux-native"], "windows")
    assert not mism.commands and len(mism.skipped) > 50            # 플랫폼 불일치는 SKIPPED 로 예고
    ora = dryrun.plan(pack, pack.profiles["oracle-db"], "linux")
    assert ora.bundles and ora.bundles[0][0] == "oracle-db"
    text = dryrun.render([lin, ora], hosts_by_platform={"linux": 3})
    assert "[linux] 호스트 3대" in text and "예상 원격 잔류물: 없음" in text and "임시 디렉터리" in text
    assert "변경 명령 0개" in text


# ------------------------------------------------------------------ preflight
class _Conn(Connection):
    def __init__(self, outputs: dict[str, str], *, privileged: bool | None = False, os: str = "linux"):
        self.outputs, self.priv, self.os_ = outputs, privileged, os
        self.cmds: list[str] = []

    def connect(self): ...
    def probe(self): return RemoteEnvironment(os=self.os_, privileged=self.priv)
    def exec(self, argv, **kw):
        cmd = argv[-1]
        self.cmds.append(cmd)
        for k, v in self.outputs.items():
            if k in cmd:
                return ExecResult(argv, 0, v, "", 1)
        return ExecResult(argv, 1, "", "", 1)
    def upload(self, l, r): ...
    def download(self, r, l): ...
    def cleanup(self, p): return CleanupReport()
    def close(self): ...


def test_preflight_reports_sqlplus_privilege_tmp_skew_and_params() -> None:
    pack = loader.load(PACK)
    conn = _Conn({"date +%s": "1000", "df -kP": "1024"})          # sqlplus 없음, sudo 안 됨, /tmp 1MB, 시계 1970년
    env = RemoteEnvironment(os="linux", privileged=False)
    notes = preflight.run(conn, env, native=pack.profiles["oracle-native"].native, has_bundles=True,
                          missing_params=["ORACLE_HOME"])
    joined = "\n".join(notes)
    assert "ORACLE_HOME 미지정" in joined and "sqlplus 없음" in joined and "일반 계정" in joined
    assert "/tmp 여유 1MB" in joined and "시간 편차" in joined
    # 전부 읽기전용 명령이었는가
    from infraguard.rules.policy import check
    assert all(check("sh", c) == [] for c in conn.cmds), conn.cmds

    fine = _Conn({"date +%s": str(int(__import__("time").time())), "command -v sqlplus": "/usr/bin/sqlplus",
                  "sudo -n": "OK"}, privileged=False)
    notes = preflight.run(fine, env, native=pack.profiles["oracle-native"].native, has_bundles=False)
    assert notes == ["일반 계정(sudo -n 가능)"]
    assert preflight.run(_Conn({}, os="windows"), RemoteEnvironment(os="windows"), native=[], has_bundles=False) == []


def test_scan_host_runs_preflight_and_keeps_going(tmp_path: Path) -> None:
    conn = _Conn({"date +%s": "1000"}, privileged=True)
    job = HostJob(native=["U-01"], preflight=True)
    host = scan_host(conn, job, tmp_path, host_id="h", hostname="h")
    assert any("시간 편차" in n for n in host.preflight)
    assert host.results and host.error is None
    off = scan_host(_Conn({"date +%s": "1000"}), HostJob(native=["U-01"], preflight=False), tmp_path, host_id="h", hostname="h")
    assert off.preflight == []
