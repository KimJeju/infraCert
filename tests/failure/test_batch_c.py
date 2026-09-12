"""배치 C — 세션 패키지(크리덴셜 없음·경로 화이트리스트)·자산 메타 스냅샷·위험도·예외 승인·수동확인 템플릿."""

from __future__ import annotations

import json
import zipfile
from datetime import date, datetime
from pathlib import Path

import pytest

from infraguard.assets.exceptions import APPROVED, EXPIRED, RiskException
from infraguard.assets.models import Host
from infraguard.assets.store import AssetStore
from infraguard.core.models import CheckResult, HostResult, ScanResult
from infraguard.core.status import Severity, Status
from infraguard.orchestrator.results_store import ResultsStore
from infraguard.reporting import html as html_report
from infraguard.reporting import xlsx as xlsx_report
from infraguard.result import gate, risk
from infraguard.workspace import package


def _stores(tmp_path: Path) -> tuple[AssetStore, ResultsStore]:
    return AssetStore(tmp_path / "assets.db"), ResultsStore(tmp_path / "results.db")


def _scan(sid: str = "s1") -> ScanResult:
    h = HostResult(host_id="h1", hostname="web01", asset={"criticality": "CRITICAL", "environment": "PROD"},
                   results=[CheckResult(rule_id="U-01", name="root", status=Status.FAIL, reason="t",
                                        severity=Severity.HIGH, evidence="PermitRootLogin yes"),
                            CheckResult(rule_id="U-02", name="pw", status=Status.FAIL, reason="t",
                                        severity=Severity.LOW, evidence="x"),
                            CheckResult(rule_id="U-03", name="ok", status=Status.PASS, reason="t")])
    return ScanResult(scan_id=sid, engine_version="0.1", rule_pack_version="kisa-2026 1", rule_pack_sha256="ab" * 32,
                      started_at=datetime(2026, 9, 12, 10, 0), finished_at=datetime(2026, 9, 12, 10, 5), hosts=[h])


# ------------------------------------------------------------------ 세션 패키지
def test_session_roundtrip_without_credentials(tmp_path: Path) -> None:
    assets, results = _stores(tmp_path / "a")
    assets.upsert(Host(host_id="h1", name="web01", address="10.0.0.1", username="igtest", environment="PROD",
                       criticality="CRITICAL", owner="김담당", tags=["dmz", "was"]))
    assets.set_exception(RiskException(host_id="h1", rule_id="U-02", reason="보상통제", expires_at="2099-01-01"))
    results.save_scan(_scan())
    logs = tmp_path / "a" / "logs"
    logs.mkdir()
    (logs / "terminal_x.log").write_text("[IN] ls\n[OUT] ****\n", encoding="utf-8")
    (logs / "crash.bin").write_bytes(b"\x00")               # .log/.txt 아니면 제외

    out = package.export_session(tmp_path / "s.zip", assets=assets, results=results, engine_version="0.1",
                                 rulepack={"name": "kisa-2026", "sha256": "ab" * 32}, audit_dir=logs)
    with zipfile.ZipFile(out) as zf:
        names = set(zf.namelist())
        assert names == {"session.json", "assets.json", "exceptions.json", "results/s1.json", "audit/terminal_x.log"}
        blob = b"".join(zf.read(n) for n in names).lower()
        # 크리덴셜 키 자체가 없다(auth_kind 값 "password" 는 방식 표시일 뿐)
        assert b'"password":' not in blob and b"passphrase" not in blob and b"secret" not in blob and b"infraguard#" not in blob
        assert json.loads(zf.read("session.json"))["rulepack"]["sha256"] == "ab" * 32

    assets2, results2 = _stores(tmp_path / "b")
    s = package.import_session(out, assets=assets2, results=results2, audit_dir=tmp_path / "b" / "logs")
    assert (s.hosts, s.scans, s.exceptions, s.audit_files, s.skipped) == (1, 1, 1, 1, [])
    h = assets2.get("h1")
    assert h and h.criticality == "CRITICAL" and h.tags == ["dmz", "was"]
    sc = results2.load_scan("s1")
    assert sc and sc.rule_pack_sha256 == "ab" * 32 and sc.hosts[0].asset["criticality"] == "CRITICAL"
    assert assets2.get_exception("h1", "U-02").reason == "보상통제"
    assert (tmp_path / "b" / "logs" / "terminal_x.log").exists()


def test_import_ignores_members_outside_whitelist(tmp_path: Path) -> None:
    z = tmp_path / "evil.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("session.json", json.dumps({"format": 1}))
        zf.writestr("../../evil.txt", "x")
        zf.writestr("audit/../escape.log", "x")
        zf.writestr("results/../../x.json", "{}")
        zf.writestr("C:/abs.json", "{}")
    assets, results = _stores(tmp_path)
    s = package.import_session(z, assets=assets, results=results, audit_dir=tmp_path / "logs")
    assert s.scans == 0 and s.audit_files == 0 and len(s.skipped) == 4
    assert not (tmp_path / "evil.txt").exists() and not (tmp_path / "logs").exists()
    with zipfile.ZipFile(tmp_path / "notpkg.zip", "w") as zf:
        zf.writestr("readme.txt", "x")
    with pytest.raises(ValueError, match="session.json"):
        package.import_session(tmp_path / "notpkg.zip", assets=assets, results=results)


# ------------------------------------------------------------------ 자산 메타 · 위험도
def test_asset_snapshot_and_risk_matrix() -> None:
    h = Host(host_id="h", name="n", address="a", environment="PROD", criticality="LOW", owner="o", role="WEB")
    assert h.asset_snapshot() == {"environment": "PROD", "criticality": "LOW", "owner": "o", "role": "WEB",
                                  "project": "기본", "group": "서버"}
    assert risk.level(Severity.HIGH, "CRITICAL") == "CRITICAL"
    assert risk.level(Severity.HIGH, "HIGH") == "CRITICAL"
    assert risk.level(Severity.MEDIUM, "HIGH") == "HIGH"
    assert risk.level(Severity.HIGH, "MEDIUM") == "HIGH"
    assert risk.level(Severity.MEDIUM, "MEDIUM") == "MEDIUM"
    assert risk.level(Severity.LOW, "CRITICAL") == "MEDIUM"
    assert risk.level(Severity.LOW, "LOW") == "LOW"
    assert risk.level(Severity.MEDIUM, "LOW") == "LOW"
    assert risk.level(None, None) == "MEDIUM"              # 모르면 중간, 추측해 올리지 않는다
    assert risk.summary(_scan()) == {"CRITICAL": 1, "HIGH": 0, "MEDIUM": 1, "LOW": 0}


# ------------------------------------------------------------------ 예외 승인
def test_exception_state_expiry_and_gate(tmp_path: Path) -> None:
    today = date(2026, 9, 12)
    e = RiskException(host_id="h1", rule_id="U-01", reason="r", expires_at="2026-12-31")
    assert e.state(today) == APPROVED and e.days_left(today) == 110 and e.label(today) == "승인 D-110"
    x = RiskException(host_id="h1", rule_id="U-02", reason="r", expires_at="2026-09-01")
    assert x.state(today) == EXPIRED and x.label(today) == "만료 11일 경과"
    assert RiskException(host_id="h", rule_id="r", reason="r").label(today) == "승인 (무기한)"

    st = AssetStore(tmp_path / "a.db")
    st.set_exception(e)
    st.set_exception(x)
    assert set(st.exception_map()) == {("h1", "U-01"), ("h1", "U-02")}
    st.remove_exception("h1", "U-01")
    assert st.get_exception("h1", "U-01") is None and len(st.all_exceptions()) == 1

    # 게이트: 만료된 예외가 취약 항목에 걸려 있으면 ✗. 판정 자체는 그대로 취약이다.
    scan = _scan()
    items = gate.run(scan, exceptions={("h1", "U-02"): RiskException(host_id="h1", rule_id="U-02", reason="r",
                                                                      expires_at="2000-01-01")})
    by = {i.label: i for i in items}
    assert not by["만료된 예외 없음"].ok and "web01/U-02" in by["만료된 예외 없음"].detail
    assert scan.hosts[0].results[1].status is Status.FAIL


def test_reports_carry_risk_and_exception_columns(tmp_path: Path) -> None:
    scan = _scan()
    exc = {("h1", "U-02"): RiskException(host_id="h1", rule_id="U-02", reason="AD 중앙관리", control="MFA+PAM",
                                         expires_at="2099-01-01")}
    h = html_report.build(scan, tmp_path / "r.html", exceptions=exc).read_text(encoding="utf-8")
    assert "<th>위험도</th>" in h and "심각" in h and "AD 중앙관리" in h
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_report.build(scan, tmp_path / "r.xlsx", exceptions=exc))
    rows = list(wb["결과"].iter_rows(min_row=2, values_only=True))
    assert rows[0][14] == "심각" and rows[1][14] == "중간" and "MFA+PAM" in rows[1][15] and rows[2][14] in ("", None)


# ------------------------------------------------------------------ 수동확인 템플릿
def test_manual_note_templates_roundtrip() -> None:
    from infraguard.ui.pages.manual import compose_note, split_note
    n = compose_note(["관리자 인터뷰", "예외 승인 확인"], "  업무상 필요, 승인문서 확인  ")
    assert n == "[관리자 인터뷰, 예외 승인 확인] 업무상 필요, 승인문서 확인"
    assert split_note(n) == (["관리자 인터뷰", "예외 승인 확인"], "업무상 필요, 승인문서 확인")
    assert split_note("[임의] 텍스트") == ([], "[임의] 텍스트")     # 템플릿 밖은 건드리지 않는다
    assert compose_note([], "x") == "x" and split_note("") == ([], "")
