"""XLSX 리포트.

시트 구성:
    요약      호스트별 · 판정별 집계 (COUNTIFS 수식 — 결과 시트를 편집해도 살아있다)
    결과      전체 항목 (호스트 x 항목)
    수동확인   UNKNOWN 만 추린 작업용 시트. 판정 입력 컬럼 포함
    실행이슈   ERROR / SKIPPED 및 정리 실패
    환경      호스트별 수집 환경
    원본추적   어느 산출물의 몇 번째 줄에서 왔는지
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from infraguard.assets.exceptions import RiskException
from infraguard.core.models import ScanResult
from infraguard.core.status import DISPLAY_KO, ORDER, Status
from infraguard.result import diff as _diff
from infraguard.result import risk as _risk
from infraguard.result.engine import provenance_text

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="D9D9D9")
STATUS_FILL = {
    Status.PASS: PatternFill("solid", fgColor="E2EFDA"),
    Status.FAIL: PatternFill("solid", fgColor="FCE4E4"),
    Status.UNKNOWN: PatternFill("solid", fgColor="FFF2CC"),
    Status.SKIPPED: PatternFill("solid", fgColor="F2F2F2"),
    Status.ERROR: PatternFill("solid", fgColor="E4DFEC"),
}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RESULT_HEADERS = ["호스트", "항목코드", "점검항목", "중요도", "진단결과",
                  "판정근거", "점검내용", "판정출처", "분석자메모", "조치방법", "판단기준", "판정추적",
                  "이전결과", "변화", "위험도", "예외"]


def _style_header(ws: Worksheet, ncols: int, row: int = 1) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.fill = HDR_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 24


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _body_font(ws: Worksheet, first_row: int = 2) -> None:
    for row in ws.iter_rows(min_row=first_row):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build(scan: ScanResult, out: Path, remediation: dict[str, str] | None = None,
          criteria: dict[str, str] | None = None, *, baseline: ScanResult | None = None,
          exceptions: dict[tuple[str, str], RiskException] | None = None) -> Path:
    """remediation/criteria: rule id → 조치방법/판단기준(룰팩 메타·가이드). 없으면 빈 컬럼 — 추측해 채우지 않는다.
    baseline: 전회 진단. 있으면 결과 시트에 이전결과·변화 컬럼, 요약 시트에 조치 현황."""
    wb = Workbook()

    ws_res = wb.active
    ws_res.title = "결과"
    d = _diff.diff(scan, baseline)
    _write_results(ws_res, scan, remediation or {}, criteria or {}, d, exceptions or {})

    _write_summary(wb.create_sheet("요약", 0), scan, ws_res.title, baseline=baseline, dsum=_diff.summary(d))
    _write_manual(wb.create_sheet("수동확인"), scan)
    _write_issues(wb.create_sheet("실행이슈"), scan)
    _write_env(wb.create_sheet("환경"), scan)
    _write_trace(wb.create_sheet("원본추적"), scan)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ------------------------------------------------------------------- 결과
def _write_results(ws: Worksheet, scan: ScanResult, remediation: dict[str, str], criteria: dict[str, str],
                   d: dict | None = None, exceptions: dict | None = None) -> None:
    d = d or {}
    exceptions = exceptions or {}
    ws.append(RESULT_HEADERS)
    for h in scan.hosts:
        for r in h.results:
            ws.append([
                h.hostname,
                r.rule_id,
                r.name,
                r.severity.value if r.severity else "",
                DISPLAY_KO[r.status],
                r.reason,
                r.evidence or "",
                "스크립트" if r.verdict_source == "script" else "분석자",
                r.analyst_note or "",
                remediation.get(r.rule_id, "") if r.status is Status.FAIL or r.status is Status.UNKNOWN else "",
                criteria.get(r.rule_id, ""),
                provenance_text(r),
                DISPLAY_KO[p] if (p := d.get((h.hostname, r.rule_id), (None, None))[0]) else "",
                _diff.LABEL_KO[c] if (c := d.get((h.hostname, r.rule_id), (None, None))[1]) else "",
                _risk.LABEL_KO[_risk.level(r.severity, h.asset.get("criticality"))] if r.status is Status.FAIL else "",
                (f"{e.label()} / {e.reason}" + (f" / 보상통제: {e.control}" if e.control else "")
                 if (e := exceptions.get((h.host_id, r.rule_id))) else ""),
            ])
    _style_header(ws, len(RESULT_HEADERS))
    _widths(ws, [16, 10, 34, 8, 12, 34, 56, 10, 24, 40, 40, 48, 10, 10, 8, 40])
    _body_font(ws)
    # 진단결과 컬럼 색상
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            for st, disp in DISPLAY_KO.items():
                if cell.value == disp:
                    cell.fill = STATUS_FILL[st]
                    cell.alignment = Alignment(horizontal="center", vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


# ------------------------------------------------------------------- 요약
def _write_summary(ws: Worksheet, scan: ScanResult, result_sheet: str, *, baseline: ScanResult | None = None,
                   dsum: dict[str, int] | None = None) -> None:
    ws["A1"] = "InfraGuard 진단 결과 요약"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)

    meta = [
        ("Scan ID", scan.scan_id),
        ("엔진 버전", scan.engine_version),
        ("룰팩 버전", scan.rule_pack_version or "-"),
        ("룰팩 SHA-256", scan.rule_pack_sha256 or "-"),
        ("프로파일", scan.profile or "-"),
        ("시작", scan.started_at.strftime("%Y-%m-%d %H:%M:%S")),
        ("종료", scan.finished_at.strftime("%Y-%m-%d %H:%M:%S") if scan.finished_at else "-"),
        ("대상 호스트", str(len(scan.hosts))),
    ]
    if baseline is not None and dsum is not None:
        rate = _diff.fix_rate(dsum)
        meta.append(("전회 진단(기준)", baseline.scan_id))
        meta += [(f"  {_diff.LABEL_KO[k]}", str(dsum.get(k, 0))) for k in _diff.ORDER]
        meta.append(("  조치율", f"{rate:.0%}" if rate is not None else "-"))
    rs = _risk.summary(scan)
    meta += [(f"위험도 {_risk.LABEL_KO[k]}", str(rs[k])) for k in _risk.LEVELS]
    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=True, size=10)
        ws.cell(row=i, column=2, value=v).font = Font(name=FONT, size=10)

    start = 3 + len(meta) + 1
    headers = ["호스트"] + [DISPLAY_KO[s] for s in ORDER] + ["합계"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=start, column=c, value=h)
    _style_header(ws, len(headers), row=start)

    # 집계는 하드코딩하지 않고 COUNTIFS 로 건다.
    # 결과 시트를 사람이 수정해도 요약이 따라간다.
    n = len(RESULT_HEADERS)
    last = 1 + sum(len(h.results) for h in scan.hosts)
    host_col = f"'{result_sheet}'!$A$2:$A${last}"
    stat_col = f"'{result_sheet}'!$E$2:$E${last}"

    for i, h in enumerate(scan.hosts, start=1):
        r = start + i
        ws.cell(row=r, column=1, value=h.hostname)
        for j in range(2, len(ORDER) + 2):
            col = get_column_letter(j)
            ws.cell(
                row=r, column=j,
                value=f'=COUNTIFS({host_col},$A{r},{stat_col},{col}${start})',
            )
        ws.cell(row=r, column=len(headers),
                value=f"=SUM(B{r}:{get_column_letter(len(ORDER) + 1)}{r})")

    total_row = start + len(scan.hosts) + 1
    ws.cell(row=total_row, column=1, value="합계").font = Font(name=FONT, bold=True, size=10)
    for j in range(2, len(headers) + 1):
        col = get_column_letter(j)
        ws.cell(row=total_row, column=j,
                value=f"=SUM({col}{start + 1}:{col}{total_row - 1})").font = Font(
            name=FONT, bold=True, size=10)

    for row in ws.iter_rows(min_row=start + 1, max_row=total_row, max_col=len(headers)):
        for cell in row:
            if cell.font is None or not cell.font.bold:
                cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            if cell.column > 1:
                cell.alignment = Alignment(horizontal="center")

    note = total_row + 2
    ws.cell(row=note, column=1,
            value="※ 집계는 '결과' 시트를 참조하는 수식입니다. 결과를 수정하면 자동 반영됩니다.")
    ws.cell(row=note, column=1).font = Font(name=FONT, size=9, italic=True)
    ws.cell(row=note + 1, column=1,
            value="※ '수동확인'은 자동 판정이 불가한 항목입니다. 실행오류(ERROR)와 구분됩니다.")
    ws.cell(row=note + 1, column=1).font = Font(name=FONT, size=9, italic=True)
    _widths(ws, [22] + [11] * (len(ORDER) + 1))
    _ = n


# --------------------------------------------------------------- 수동확인
def _write_manual(ws: Worksheet, scan: ScanResult) -> None:
    headers = ["호스트", "항목코드", "점검항목", "중요도", "수집근거",
               "판정입력(양호/취약/해당없음)", "판정사유"]
    ws.append(headers)
    for h in scan.hosts:
        for r in h.results:
            if r.status is Status.UNKNOWN:
                ws.append([h.hostname, r.rule_id, r.name,
                           r.severity.value if r.severity else "",
                           r.evidence or r.reason, "", ""])
    _style_header(ws, len(headers))
    _widths(ws, [16, 10, 34, 8, 70, 20, 30])
    _body_font(ws)
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="FFFFCC")
    ws.freeze_panes = "C2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


# --------------------------------------------------------------- 실행이슈
def _write_issues(ws: Worksheet, scan: ScanResult) -> None:
    headers = ["호스트", "구분", "항목코드", "내용"]
    ws.append(headers)
    for h in scan.hosts:
        if h.error:
            ws.append([h.hostname, "호스트오류", "", h.error])
        if h.cleanup_ok is False:
            ws.append([h.hostname, "원격정리실패", "",
                       f"잔류: {', '.join(h.cleanup_leftovers)}"])
        for r in h.results:
            if r.status in (Status.ERROR, Status.SKIPPED):
                ws.append([h.hostname, DISPLAY_KO[r.status], r.rule_id, r.reason])
    _style_header(ws, len(headers))
    _widths(ws, [16, 14, 10, 80])
    _body_font(ws)


# ------------------------------------------------------------------- 환경
def _write_env(ws: Worksheet, scan: ScanResult) -> None:
    headers = ["호스트", "OS", "버전", "아키텍처", "계정", "권한", "셸",
               "로케일", "인코딩", "정보완전성", "비고"]
    ws.append(headers)
    for h in scan.hosts:
        e = h.environment
        ws.append([
            h.hostname, e.os or "", e.os_version or "", e.architecture or "",
            e.user or "", "root" if e.privileged else ("일반" if e.privileged is False else ""),
            e.shell or "", e.locale or "", e.encoding or "",
            "불완전" if e.incomplete else "완전", "; ".join(e.notes),
        ])
    _style_header(ws, len(headers))
    _widths(ws, [16, 10, 26, 10, 16, 8, 14, 16, 10, 10, 40])
    _body_font(ws)


# --------------------------------------------------------------- 원본추적
def _write_trace(ws: Worksheet, scan: ScanResult) -> None:
    headers = ["호스트", "항목코드", "번들", "산출물", "파서프로파일", "줄번호", "경고"]
    ws.append(headers)
    for h in scan.hosts:
        for r in h.results:
            s = r.source
            ws.append([h.hostname, r.rule_id, s.bundle_id or "", s.artifact or "",
                       s.profile or "", s.line or "", "; ".join(r.warnings)])
    _style_header(ws, len(headers))
    _widths(ws, [16, 10, 20, 28, 18, 8, 48])
    _body_font(ws)
